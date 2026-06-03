"""
Full correction pass on dashboard_מלא_עם_ביקורת_4.xlsx:
  1. Sort each Q sheet (level asc, ללא AI before עם AI)
  2. Apply rubric-based level corrections (existing rules)
  3. Resolve all ⚠=default uncertainty markers → definitive classification
  4. Rebuild summary tables
Output: dashboard_סופי_מתוקן.xlsx
"""

import re
from collections import defaultdict
import openpyxl
from openpyxl.cell.cell import MergedCell

INPUT  = "/root/.claude/uploads/e68b10a7-683f-414f-924b-e35511b15693/298b66ae-dashboard_______________4.xlsx"
OUTPUT = "/home/user/hany/dashboard_סופי_מתוקן.xlsx"

LEVEL_LABELS = {
    1: {1:'רמה 1 – הימנעות / חוסר אסטרטגיה', 2:'רמה 2 – פסיביות / תלות',
        3:'רמה 3 – פרואקטיביות בסיסית', 4:'רמה 4 – ניהול אסטרטגי'},
    2: {1:'רמה 1 – ויתור ותגובה רגשית קיצונית', 2:'רמה 2 – תלות מיידית במבוגר',
        3:'רמה 3 – עצמאות בסיסית וחיפוש מידע', 4:'רמה 4 – ויסות עצמי מתוכנן'},
    3: {1:'רמה 1 – ייחוס חיצוני ובלתי נשלט', 2:'רמה 2 – ייחוס פנימי כללי ובלתי ספציפי',
        3:'רמה 3 – זיהוי פגם אסטרטגי ספציפי', 4:'רמה 4 – ניתוח עומק קוגניטיבי'},
    4: {1:'רמה 1 – הערכה חיצונית בלעדית', 2:'רמה 2 – הערכה חיצונית עם הסתכלות על טעויות',
        3:'רמה 3 – הערכה פנימית ומכוונת תהליך', 4:'רמה 4 – הערכה עצמית אסטרטגית'},
    5: {1:'רמה 1 – חוסר עניין ורוגז', 2:'רמה 2 – נושאים מוכתבים מהסילבוס',
        3:'רמה 3 – סקרנות ממוקדת בנושאים', 4:'רמה 4 – למידה יישומית ופרקטית'},
    6: {1:'רמה 1 – תגובה רגשית הרסנית/קיצונית', 2:'רמה 2 – רגש שלילי פסיבי',
        3:'רמה 3 – השלמה וקבלה', 4:'רמה 4 – אכזבה בונה ורצון לשיפור'},
    7: {1:'רמה 1 – הימנעות ובושה', 2:'רמה 2 – ציות פסיבי',
        3:'רמה 3 – סקרנות בסיסית', 4:'רמה 4 – חיפוש עזרה אקטיבי'},
    8: {1:'רמה 1 – חוסר מטרות', 2:'רמה 2 – שאלה טכנית/ציון',
        3:'רמה 3 – אבחון טעויות', 4:"רמה 4 – בקשת שיפור ומועד ב'"},
    9: {1:'רמה 1 – דחיינות מוחלטת', 2:'רמה 2 – תכנון ימים ספורים מראש',
        3:'רמה 3 – תכנון שבועי', 4:'רמה 4 – תכנון אסטרטגי ל-2 שבועות'},
}


def get_level_num(val):
    if val is None:
        return None
    m = re.search(r'רמה\s*(\d)', str(val))
    return int(m.group(1)) if m else None


def safe_write(ws, row, col, value):
    cell = ws.cell(row, col)
    if not isinstance(cell, MergedCell):
        cell.value = value


# ─── Rubric corrections (confirmed cases only) ──────────────────────────────
def apply_correction(q, lv, ans):
    """Returns (new_level, reason) or (None, None)."""
    if ans is None:
        return None, None
    a = str(ans).strip()

    if q == 1:
        if lv == 2 and re.match(
            r'^(chat\s*gpt|chatgpt|צ[׳\'״]?אט\s*$|^gpt$|ג[׳\'״]?יפיטי|'
            r'^ai$|^בינה$|gemini|copilot|ישתמש\s*בצאט|אפתח\s*gpt|'
            r'^גוגל$|^יוטיוב$)[\s.,]*$', a, re.IGNORECASE):
            return 1, 'L1: כלי בודד ללא הסבר כיצד משתמשים'
        if lv == 2 and len(a) <= 8 and not re.search(
                r'(AI|בינה|צאט|גוגל|יוטיוב|אבא|אמא|חבר|מורה)', a):
            return 1, 'L1: תגובה ≤8 תווים ללא כוונה פעילה'
        if lv == 2 and len(a) > 8:
            person = re.search(r'(אבא|אמא|הורים|חבר|חברה|מורה|משפחה|אחי|אחות|סבא|סבתא)', a)
            indep  = re.search(r'(גוגל|יוטיוב|אינטרנט|AI|ai|בינה|צ[׳\'״]?אט|gpt|'
                               r'ג[׳\'״]?יפיטי|חפש|אחפש|סרטון)', a, re.IGNORECASE)
            if indep and not person:
                return 3, 'L3: AI/חיפוש עצמאי ללא הסתמכות על אדם'
        if lv == 3 and re.search(r'(אבחן|בחן|חידון|שאלון|אוודא|תרגול|אתרגל|בדוק|בודק)', a):
            return 4, 'L4: חיפוש עצמאי + בחינה עצמית'

    elif q == 2:
        if lv == 2 and len(a) <= 12 and not re.search(r'(אנסה|ינסה|אבדוק|אחפש|קרא|יוטיוב|גוגל)', a):
            return 1, 'L1: פנייה מיידית ללא ניסיון עצמי'
        if lv == 2:
            self_try = re.search(r'(קרא\s*שוב|חפש|יוטיוב|גוגל|אנסה|ינסה|שוב|מחדש)', a)
            person   = re.search(r'(אבא|אמא|הורים|חבר|מורה|משפחה)', a)
            after    = re.search(r'(אחרי|אם\s*לא|אם\s*עדיין|ולאחר\s*מכן|רק\s*אחר)', a)
            if self_try and person and after:
                return 3, 'L3: ניסיון עצמי → פנייה לאדם אחרי כישלון'

    elif q == 3:
        if lv == 1 and re.search(r'(לא\s*למדתי|לא\s*השקעתי|לא\s*הכנתי|לא\s*הבנתי|'
                                  r'לא\s*תרגלתי|לא\s*ידעתי|לא\s*קראתי)', a):
            return 2, 'L2: ייחוס פנימי כללי – לא למדתי מספיק'
        if lv == 2 and re.match(r'^(הלחץ|בלאק[\s\-]?אאוט|חירטטתי|ככה\s*ככה|'
                                 r'בלבול|לחץ|מהלחץ|הייתי\s*בלחץ|בגלל\s*לחץ)[\s.,]*$', a.strip()):
            return 1, 'L1: ייחוס חיצוני בלבד – לחץ ללא לקיחת אחריות'
        if lv == 2 and re.search(r'(שיטת\s*הלמידה|דרך\s*לא\s*נכונה|לא\s*תרגלתי|'
                                  r'שאלות\s*מילוליות|לא\s*חזרתי|הזנחתי|'
                                  r'טעויות\s*קטנות\s*בחישוב|לא\s*קראתי\s*את\s*השאלה)', a):
            return 3, 'L3: זיהוי מנגנון ספציפי של הכשל'

    elif q == 4:
        if lv == 2 and re.search(r'(היה\s*לי\s*קל|הרגשתי|הצלחתי\s*לענות|בזמן\s*המבחן|היה\s*קל)', a):
            if not re.search(r'(לפי\s*הציון\s*בלבד|הציון\s*הוא)', a):
                return 3, 'L3: מדד פנימי סובייקטיבי'

    elif q == 5:
        if lv == 3 and re.search(
                r'(אותו\s*דבר|כמו\s*ש(למדתי|עשיתי|הבנתי|היה|הצלחתי)|'
                r'כמו\s*(ב)?פעם|כמו\s*(ב)?חומר\s*הקוד|'
                r'ארצה\s*ש(החומר|זה)\s*(יהיה)?\s*דומ|'
                r'^ללמוד\s*(ולחזור\s*)?(על\s*החומר)?\s*$)', a):
            return 2, 'L2: חזרה על אסטרטגיה קודמת ללא סקרנות ממוקדת'
        if lv == 2 and re.search(
                r'(אנרגי|פיזיק|כימי|ביולוג|גוף\s*(ה)?אדם|חשמל|מגנ|'
                r'מחשב|תכנות|היסטורי|גיאוגרפי|ספרות|פילוסופ|'
                r'מתמטיק|סטטיסטיק|גיאומטרי|אסטרונ|חלל|רפואה|'
                r'מוסיק|אמנות|פסיכולוג|כלכל|חינוך|חברה)', a, re.IGNORECASE
        ) and not re.search(r'(אותו\s*דבר|כמו\s*ש)', a):
            return 3, 'L3: ציון נושא ספציפי – סקרנות ממוקדת'
        if lv == 3 and re.search(r'(ניסוי|לחקור|חקירה|מחקר|יישום|להתנסות|לעשות\s*בעצמ)', a):
            return 4, 'L4: למידה פרקטית – ניסויים / חקירה עצמית'

    elif q == 6:
        action = (r'(אשתפר|ישפר|לשפר|ללמוד|אלמד|אנסה|ינסה|אשתדל|לנסות|'
                  r'אבוא|למבחן|יותר\s*טוב|השתפר|אשנה|אשקיע|ישקיע|'
                  r'לפעם\s*הבאה|הבאה|נסה|אטרח|אוסיף|אדבר)')
        if lv == 2 and len(a) <= 40 and not re.search(action, a):
            return 1, 'L1: תגובה רגשית ללא כוונת פעולה'

    elif q == 7:
        if lv == 1 and re.search(r'^כן', a) and not re.search(
                r'(לא\s*רוצ|לא\s*ארצ|לא\s*אוהב|מביך|בושה|מפחיד|לא\s*נוח|לא\s*ארצה\s*לדבר)', a):
            return 2, 'L2: ציות פסיבי – הולך ללא יוזמה'
        if lv == 2 and re.search(
                r'(מבחן\s*חוזר|מועד\s*ב|שיפור\s*ציון|כלים\s*(ל|ש|כד)|'
                r'איך\s*(ל|א)(השתפר|שפר|הצליח)|כדי\s*להשתפר|'
                r'מה\s*(עוד\s*)?(אפשר|כדאי|אני\s*(צריך|יכול))\s*(לעשות|לשפר))', a):
            return 4, 'L4: חיפוש עזרה אקטיבי – כלים / מועד ב\' / שיפור מכוון'
        if lv == 2 and re.search(r'(במה\s*טעיתי|הטעויות\s*שלי|מה\s*(לא\s*)?עשיתי\s*(לא\s*)?נכון)', a):
            if not re.search(r'(ואיך|ולשפר|כדי\s*להשתפר|שיפור)', a):
                return 3, 'L3: שאלה על טעויות ספציפיות'
        if lv == 3 and re.search(
                r'(במה\s*טעיתי|הטעויות\s*שלי|מה\s*(לא\s*)?עשיתי\s*(לא\s*)?נכון)', a
        ) and re.search(r'(ואיך|ולשפר|כדי\s*להשתפר|שיפור|פעם\s*הבאה|לפעם)', a):
            return 4, 'L4: טעויות + כוונת שיפור'

    elif q == 8:
        if lv == 1 and re.search(r'(מורה|הורים|אבא|אמא|חבר\b|חברה\b|אפנה|יפנה)', a) \
                and not re.search(r'(לא\s*(יודע|רוצ)|כלום|שום)', a):
            return 2, 'L2: פנייה לגורם עזרה חיצוני'
        if lv == 2 and re.search(
                r'(להבין\s*(מה|איפה|למה|קודם)|מה\s*לא\s*(הבנתי|ידעתי)|'
                r'אאבחן|אזהה|לאתר\s*את\s*הטעות)', a):
            return 3, 'L3: אבחון עצמי של פערי הבנה'
        if lv == 2 and re.search(r'(מועד\s*ב|מבחן\s*חוזר)', a):
            return 4, "L4: בקשת מועד ב'"

    elif q == 9:
        if lv == 2 and len(a) <= 15 and not re.search(r'(שבוע|יום|שעה|לפני|מראש)', a):
            return 1, 'L1: הצהרה כללית ללא תוכנית זמן'
        if lv == 2 and re.search(r'(שבוע|שבועיים|יום|שעה|מראש|מוקדם)', a) and re.search(
                r'(אדע\s*שמוכן|ארגיש|מבחן\s*לדוגמ|אבחן|שאלון|תרגילים)', a):
            return 3, 'L3: לוח זמנים ספציפי + מדד מוכנות'

    return None, None


# ─── Resolve ⚠ uncertainty → definitive classification ────────────────────
def resolve_uncertain(q, lv, ans):
    """
    Always returns (level, clean_reason).
    Resolves the ⚠=default uncertainty marker based on rubric.
    """
    if ans is None:
        return lv, _clean_reason_for_level(q, lv)
    a = str(ans).strip()

    if q == 1:
        # Short/empty → L1
        if len(a) <= 4 or re.match(r'^(לא\s*יודע|אין\s*לי\s*מושג|כלום|לא|לא\s*הבנתי)[\s.,]*$', a, re.IGNORECASE):
            return 1, 'L1: חוסר כוונה / אין תשובה'
        # Person + AI/search (mixed) → L2
        person = re.search(r'(אבא|אמא|הורים|חבר|חברה|מורה|משפחה|אחי|אחות|סבא|סבתא)', a)
        indep  = re.search(r'(גוגל|יוטיוב|אינטרנט|AI|ai|בינה|צ[׳\'״]?אט|gpt|'
                           r'ג[׳\'״]?יפיטי|חפש|אחפש|סרטון)', a, re.IGNORECASE)
        # AI/search without person → L3
        if indep and not person:
            return 3, 'L3: חיפוש עצמאי – AI / מקורות מידע ללא הסתמכות על אדם'
        # Self-test planned → L4
        if re.search(r'(אבחן|בחן|חידון|שאלון|אוודא|תרגול|אתרגל|בדוק|בודק)', a):
            if indep or re.search(r'(גוגל|יוטיוב|אינטרנט|חפש|אחפש)', a, re.IGNORECASE):
                return 4, 'L4: חיפוש + בחינה עצמית'
        # Default → L2
        return 2, 'L2: אסטרטגיה כללית – תלות בסיסית או כוונה לא ממוקדת'

    elif q == 2:
        if len(a) <= 4 or re.match(r'^(לא\s*יודע|אין\s*לי\s*מושג|כלום|אוותר|'
                                    r'לא\s*אעשה|ינסה\s*לזכור)[\s.,]*$', a, re.IGNORECASE):
            return 1, 'L1: ויתור – חוסר אסטרטגיה'
        person   = re.search(r'(אבא|אמא|הורים|חבר|חברה|מורה|משפחה)', a)
        self_try = re.search(r'(קרא\s*שוב|חפש|יוטיוב|גוגל|אנסה|ינסה|שוב|מחדש|'
                             r'AI|ai|בינה|צ[׳\'״]?אט|gpt)', a, re.IGNORECASE)
        after    = re.search(r'(אחרי|אם\s*לא|אם\s*עדיין|ולאחר\s*מכן|רק\s*אחר)', a)
        if self_try and person and after:
            return 3, 'L3: ניסיון עצמי ראשון → פנייה לאדם אחרי כישלון'
        if self_try and not person:
            return 3, 'L3: חיפוש עצמאי – AI / מקורות ללא תלות מיידית'
        return 2, 'L2: פנייה מיידית לאדם או תלות בסיסית'

    elif q == 3:
        if re.match(r'^(לא\s*יודע|אין\s*לי\s*מושג|כלום)[\s.,]*$', a):
            return 1, 'L1: הימנעות מייחוס'
        if re.search(r'(המורה\s*לא\s*טוב|הסביבה|הרעש|הכיתה|הבחינה\s*הייתה\s*קשה|'
                     r'לחץ|בלאק[\s\-]?אאוט|מזל|לא\s*יודע\s*למה)', a):
            if not re.search(r'(לא\s*למדתי|לא\s*השקעתי|לא\s*הכנתי)', a):
                return 1, 'L1: ייחוס חיצוני'
        if re.search(r'(שיטת\s*הלמידה|דרך\s*לא\s*נכונה|לא\s*תרגלתי|שאלות\s*מילוליות|'
                     r'לא\s*חזרתי|הזנחתי|טעויות\s*קטנות|לא\s*קראתי\s*את\s*השאלה)', a):
            return 3, 'L3: זיהוי מנגנון ספציפי'
        return 2, 'L2: ייחוס פנימי כללי'

    elif q == 4:
        # Internal feeling → L3
        if re.search(r'(מרגיש|מרגישה|ידעתי\s*(ש|את)|הרגשתי|ידעה?\s*(ש|את)|'
                     r'קל\s*לי|בטוח\s*(ש|ב)|בטוחה\s*(ש|ב))', a):
            return 3, 'L3: מדד פנימי – תחושה / ביטחון עצמי כמדד'
        # Checking specific content / errors → L2
        if re.search(r'(טעויות|מה\s*השגתי|מה\s*לא\s*הצלחתי|מה\s*השגתי|'
                     r'מה\s*ידעתי|מה\s*לא\s*ידעתי|כמה\s*שאלות|כל\s*מה\s*שהיה\s*במבחן)', a):
            return 2, 'L2: בדיקת תוצאות / טעויות כסמן'
        # Grade as only indicator → L1
        if re.search(r'(ציון|תוצאה|הצלחתי\s*במבחן|יצא\s*טוב|יצא\s*גבוה|ציון\s*גבוה)', a):
            return 1, 'L1: הציון כמדד יחיד'
        # Generic "I will learn more" → L1 (doesn't answer "how will you know")
        return 1, 'L1: לא עונה לשאלה "כיצד תדע?" – מתאר פעולה ולא מדד'

    elif q == 5:
        if len(a) <= 4 or re.match(r'^(לא\s*יודע|אין\s*לי\s*מושג|כלום|לא|לא\s*הבנתי)[\s.,]*$', a, re.IGNORECASE):
            return 1, 'L1: חוסר עניין / אין תשובה'
        # Specific topic of interest → L3
        if re.search(r'(אנרגי|פיזיק|כימי|ביולוג|גוף\s*(ה)?אדם|חשמל|מגנ|'
                     r'מחשב|תכנות|היסטורי|גיאוגרפי|ספרות|פילוסופ|'
                     r'מתמטיק|סטטיסטיק|גיאומטרי|אסטרונ|חלל|רפואה|'
                     r'מוסיק|אמנות|פסיכולוג|כלכל|חינוך|חברה|על\s*AI|על\s*בינה|'
                     r'על\s*טכנולוג|על\s*הפלסטינ|על\s*גרעיני)', a, re.IGNORECASE
                     ) and not re.search(r'(אותו\s*דבר|כמו\s*ש)', a):
            # Practical/applied → L4
            if re.search(r'(ניסוי|לחקור|חקירה|מחקר|יישום|להתנסות|לעשות\s*בעצמ)', a):
                return 4, 'L4: למידה פרקטית ביישום ממשי'
            return 3, 'L3: ציון נושא ספציפי – סקרנות ממוקדת'
        # "Same as before" or wants AI as tool (not topic) → L2
        return 2, 'L2: כוונה כללית ללא נושא ספציפי / כלי AI כאמצעי'

    elif q == 6:
        # L1: extreme emotion OR complete indifference OR mismatch reaction
        if re.search(r'(מטורף|יפגע\s*בעצמ|לא\s*אכפת\s*לי|כיף\s*מאוד|'
                     r'^(לא\s*יודע|לא\s*יודעת|אין\s*לי\s*מושג)[\s.,]*$)', a, re.IGNORECASE):
            return 1, 'L1: תגובה לא מותאמת / אדישות / חוסר עיבוד'
        # L4: constructive – will study more next time, action plan
        if re.search(r'(פעם\s*הבאה.*לל?מוד|אלמד\s*יותר|ילמד\s*יותר|'
                     r'לפעם\s*הבאה.*טוב|אשקיע\s*יותר|ישקיע\s*יותר|'
                     r'אחזור\s*על\s*החומר|אתחיל\s*(ל|מ))', a):
            return 4, 'L4: אכזבה בונה – כוונת שיפור / תכנית פעולה'
        # L3: acceptance, equanimity, "OK", moving on without plan
        if re.search(r'(בסדר|סבבה|זה\s*חיים|אין\s*מה\s*לעשות|מקבל|'
                     r'לא\s*נורא|זה\s*בסדר|מתקדם|אמשיך)', a):
            return 3, 'L3: השלמה וקבלה – עיבוד רגשי מאוזן'
        # L2: passive negative emotion
        return 2, 'L2: רגש שלילי פסיבי – עצב / תסכול ללא כוונת שיפור'

    elif q == 7:
        # L1: explicit refusal / strong avoidance with nothing to discuss
        if re.match(r'^(לא|לא\s*ילך|לא\s*אלך|לא\s*הייתי\s*הולך)[\s.,]*$', a.strip()):
            return 1, 'L1: סירוב מפורש'
        if re.search(r'(כלום.*אין\s*לי\s*(מה|על\s*מה)|אין\s*לי.*כלום)', a):
            return 1, 'L1: הימנעות – אין מה לדון'
        # L4: asks for improvement tools / מועד ב'
        if re.search(r'(מבחן\s*חוזר|מועד\s*ב|כלים\s*(ל|ש|כד)|'
                     r'איך\s*(ל|א)(השתפר|שפר|הצליח)|כדי\s*להשתפר)', a):
            return 4, "L4: חיפוש עזרה אקטיבי – כלים / מועד ב'"
        # L3: asking specifically about mistakes
        if re.search(r'(במה\s*טעיתי|הטעויות\s*שלי|מה\s*(לא\s*)?עשיתי\s*(לא\s*)?נכון|'
                     r'מה\s*השגתי\s*לא\s*נכון)', a):
            return 3, 'L3: שאלה ממוקדת על טעויות'
        # Passive/uncertain → L2
        return 2, 'L2: ציות פסיבי – הולך ללא יוזמה ספציפית'

    elif q == 8:
        if len(a) <= 5 or re.match(r'^(לא\s*יודע|כלום|לא|לא\s*יודעת|אבכה|'
                                    r'לא\s*יאסה\s*כלום)[\s.,]*$', a, re.IGNORECASE):
            return 1, 'L1: חוסר מטרות / אין תוכנית'
        # L4: מועד ב' / asking for improvement path
        if re.search(r'(מועד\s*ב|מבחן\s*חוזר)', a):
            return 4, "L4: בקשת מועד ב'"
        # L3: structured/diagnostic approach
        if re.search(r'(להבין\s*(מה|איפה|קודם)|מה\s*לא\s*(הבנתי|ידעתי)|'
                     r'אסדר|מסדר|לפי\s*רמת\s*הקושי|ארגן|לזהות|לאתר|'
                     r'אאבחן)', a):
            return 3, 'L3: גישה מובנית / אבחון עצמי לפני פנייה'
        # Generic AI/help → L2
        return 2, 'L2: פנייה כללית לעזרה – AI / הורים / חברים ללא אבחון'

    elif q == 9:
        if len(a) <= 5 or re.match(r'^(לא\s*יודע|כלום|לא|עכשיו|ביום\s*המבחן)[\s.,]*$', a, re.IGNORECASE):
            return 1, 'L1: דחיינות / אין תוכנית'
        if re.search(r'(שבועיים|2\s*שבוע)', a):
            return 4, 'L4: תכנון אסטרטגי – שבועיים מראש'
        if re.search(r'(שבוע\s*לפני|שבוע\s*מראש|שבוע\s*קודם)', a):
            return 3, 'L3: תכנון שבועי'
        # Criterion for readiness (not time-based but mastery-based) → L3
        if re.search(r'(עד\s*ש(אצליח|אדע|אבין|אוכל\s*לפתור)|'
                     r'אדע\s*ש(מוכן|הבנתי|ידעתי)|'
                     r'עד\s*(שאני|שהמצב))', a):
            return 3, 'L3: מדד שליטה – לומד עד שמצליח'
        return 2, 'L2: תכנון כללי ללא פירוט זמן / ימים ספורים'

    return lv, _clean_reason_for_level(q, lv)


def _clean_reason_for_level(q, lv):
    """Generic clean reason when nothing specific can be determined."""
    labels = {
        1: {1:'L1: חוסר אסטרטגיה', 2:'L2: תלות / פסיביות',
            3:'L3: פרואקטיביות בסיסית', 4:'L4: ניהול אסטרטגי'},
        2: {1:'L1: ויתור', 2:'L2: תלות מיידית במבוגר',
            3:'L3: עצמאות בסיסית', 4:'L4: ויסות עצמי מתוכנן'},
        3: {1:'L1: ייחוס חיצוני', 2:'L2: ייחוס פנימי כללי',
            3:'L3: זיהוי פגם ספציפי', 4:'L4: ניתוח עומק'},
        4: {1:'L1: הציון כמדד יחיד', 2:'L2: בדיקת תוצאות',
            3:'L3: מדד פנימי', 4:'L4: הערכה עצמית אסטרטגית'},
        5: {1:'L1: חוסר עניין', 2:'L2: כוונה כללית',
            3:'L3: נושא ספציפי', 4:'L4: למידה פרקטית'},
        6: {1:'L1: תגובה לא מותאמת', 2:'L2: רגש שלילי פסיבי',
            3:'L3: השלמה וקבלה', 4:'L4: אכזבה בונה'},
        7: {1:'L1: הימנעות', 2:'L2: ציות פסיבי',
            3:'L3: שאלה על טעויות', 4:'L4: חיפוש עזרה אקטיבי'},
        8: {1:'L1: חוסר מטרות', 2:'L2: פנייה כללית',
            3:"L3: אבחון טעויות", 4:"L4: מועד ב'"},
        9: {1:'L1: דחיינות', 2:'L2: תכנון ימים ספורים',
            3:'L3: תכנון שבועי', 4:'L4: תכנון אסטרטגי'},
    }
    return labels.get(q, {}).get(lv, f'L{lv}')


# ─── Process one sheet ──────────────────────────────────────────────────────
def process_sheet(wb, q):
    ws = wb[f'Q{q}']
    labels = LEVEL_LABELS[q]
    max_col = ws.max_column

    # Locate key rows
    data_header_row = None
    summary_rows = {}
    total_row = None
    avg_row = None

    for r in range(1, 35):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        if a == 'רמה' and b == 'עם/ללא AI':
            data_header_row = r
        if (a and isinstance(a, str) and b and isinstance(b, str)
                and 'רמה' in a and get_level_num(a) is not None
                and b not in ('עם/ללא AI',)
                and isinstance(c, (int, float))):
            summary_rows[get_level_num(a)] = r
        if a == 'סה"כ' and not isinstance(ws.cell(r, 3), MergedCell):
            total_row = r
        if a == 'רמה ממוצעת':
            avg_row = r

    if data_header_row is None:
        print(f'  [שגיאה] Q{q}: data header row not found')
        return 0, 0

    # Read all data rows
    data_rows = []
    for r in range(data_header_row + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a is None or b is None:
            continue
        lv_n = get_level_num(a)
        if lv_n is None:
            continue
        if b not in ('עם AI', 'ללא AI'):
            continue
        reason = ws.cell(r, 5).value
        is_uncertain = reason and '⚠' in str(reason)
        data_rows.append([lv_n, b,
                          ws.cell(r, 3).value,  # school
                          ws.cell(r, 4).value,  # answer
                          reason,
                          is_uncertain])

    # Apply corrections and resolve uncertainty
    rubric_fixes = 0
    uncertainty_resolved = 0

    for row in data_rows:
        lv, ai, school, ans, reason, is_uncertain = row

        if is_uncertain:
            # First: try rubric correction
            new_lv, new_reason = apply_correction(q, lv, ans)
            if new_lv is not None and new_lv != lv:
                row[0] = new_lv
                row[4] = new_reason
                rubric_fixes += 1
            else:
                # Resolve uncertainty
                resolved_lv, resolved_reason = resolve_uncertain(q, lv, ans)
                row[0] = resolved_lv
                row[4] = resolved_reason
                if resolved_lv != lv:
                    rubric_fixes += 1
            uncertainty_resolved += 1
        else:
            # Non-uncertain: apply rubric correction only
            new_lv, new_reason = apply_correction(q, lv, ans)
            if new_lv is not None and new_lv != lv:
                row[0] = new_lv
                row[4] = new_reason
                rubric_fixes += 1

    # Sort: level asc, ללא AI before עם AI
    data_rows.sort(key=lambda x: (x[0], 0 if x[1] == 'ללא AI' else 1))

    # Recount
    counts = defaultdict(lambda: {'עם': 0, 'ללא': 0})
    for lv, ai, *_ in data_rows:
        counts[lv]['עם' if ai == 'עם AI' else 'ללא'] += 1

    total    = len(data_rows)
    total_ai = sum(v['עם'] for v in counts.values())
    total_no = sum(v['ללא'] for v in counts.values())
    avg      = (sum(lv * (counts[lv]['עם'] + counts[lv]['ללא'])
                    for lv in range(1, 5)) / total) if total else 0

    # Update summary table
    for lv_n in range(1, 5):
        r = summary_rows.get(lv_n)
        if r is None:
            continue
        n_ai  = counts[lv_n]['עם']
        n_no  = counts[lv_n]['ללא']
        n_tot = n_ai + n_no
        safe_write(ws, r, 3, n_ai)
        safe_write(ws, r, 4, f'{n_ai / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, r, 5, n_no)
        safe_write(ws, r, 6, f'{n_no / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, r, 7, n_tot)
        safe_write(ws, r, 8, f'{n_tot / total * 100:.1f}%' if total else '0.0%')

    if total_row:
        safe_write(ws, total_row, 3, total_ai)
        safe_write(ws, total_row, 4, f'{total_ai / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, total_row, 5, total_no)
        safe_write(ws, total_row, 6, f'{total_no / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, total_row, 7, total)
        safe_write(ws, total_row, 8, '100%')

    if avg_row:
        safe_write(ws, avg_row, 2, f'{avg:.2f}')

    # Update info line
    info_row = data_header_row - 1
    if ws.cell(info_row, 1).value and 'כל התשובות' in str(ws.cell(info_row, 1).value):
        safe_write(ws, info_row, 6, f'סה"כ: {total} | AI: {total_ai}')

    # Rewrite raw data section
    write_r = data_header_row + 1
    current_level = None

    for lv, ai, school, ans, reason, _ in data_rows:
        if lv != current_level:
            current_level = lv
            n_in_level = counts[lv]['עם'] + counts[lv]['ללא']
            safe_write(ws, write_r, 1, f'{labels[lv]} ({n_in_level} תשובות)')
            for c in range(2, max_col + 1):
                safe_write(ws, write_r, c, None)
            write_r += 1

        safe_write(ws, write_r, 1, f'רמה {lv}')
        safe_write(ws, write_r, 2, ai)
        safe_write(ws, write_r, 3, school)
        safe_write(ws, write_r, 4, ans)
        safe_write(ws, write_r, 5, reason)
        for c in range(6, max_col + 1):
            safe_write(ws, write_r, c, None)
        write_r += 1

    # Clear leftover rows
    while write_r <= ws.max_row:
        for c in range(1, max_col + 1):
            safe_write(ws, write_r, c, None)
        write_r += 1

    return rubric_fixes, uncertainty_resolved


# ─── Main ───────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(INPUT)
grand_fixes = 0
grand_resolved = 0

for q in range(1, 10):
    fixes, resolved = process_sheet(wb, q)
    grand_fixes += fixes
    grand_resolved += resolved
    print(f'  ✓ Q{q}: {fixes} תיקוני רמה, {resolved} חוסרי ודאות הוכרעו')

wb.save(OUTPUT)
print(f'\n✓ נשמר: {OUTPUT}')
print(f'סה"כ תיקוני רמה: {grand_fixes}')
print(f'סה"כ ⚠ שהוכרעו: {grand_resolved}')

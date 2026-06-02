"""
Apply rubric-based level corrections to dashboard_מיון_מתוקן.xlsx
Saves as dashboard_מתוקן_מחוון.xlsx
"""

import re
from collections import defaultdict
import openpyxl
from openpyxl.cell.cell import MergedCell

INPUT_PATH  = "/home/user/hany/dashboard_מיון_מתוקן.xlsx"
OUTPUT_PATH = "/home/user/hany/dashboard_מתוקן_מחוון.xlsx"

LEVEL_LABELS = {
    1: {1:'רמה 1 – הימנעות / חוסר אסטרטגיה', 2:'רמה 2 – פסיביות / תלות',
        3:'רמה 3 – פרואקטיביות בסיסית', 4:'רמה 4 – ניהול אסטרטגי'},
    2: {1:'רמה 1 – ויתור ותגובה רגשית קיצונית', 2:'רמה 2 – תלות מיידית במבוגר',
        3:'רמה 3 – עצמאות בסיסית וחיפוש מידע', 4:'רמה 4 – ויסות עצמי מתוכנן'},
    3: {1:'רמה 1 – ייחוס חיצוני ובלתי נשלט', 2:'רמה 2 – ייחוס פנימי כללי ובלתי ספציפי',
        3:'רמה 3 – זיהוי פגם אסטרטגי ספציפי', 4:'רמה 4 – ניתוח עומק קוגניטיבי'},
    4: {1:'רמה 1 – הערכה חיצונית בלעדית', 2:'רמה 2 – הערכה חיצונית עם הסתכלות על טעויות',
        3:'רמה 3 – הערכה פנימית ומכוונת תהליך', 4:'רמה 4 – הערכה עצמית אסטרטגית'},
    5: {1:'רמה 1 – חוסר עניין ורוגז', 2:'רמה 2 – נושאים מוכתבים מהסילבוס',
        3:'רמה 3 – סקרנות ממוקדת בנושאים', 4:'רמה 4 – למידה יישומית ופרקטית'},
    6: {1:'רמה 1 – תגובה רגשית הרסנית/קיצונית', 2:'רמה 2 – רגש שלילי פסיבי',
        3:'רמה 3 – השלמה וקבלה', 4:'רמה 4 – אכזבה בונה ורצון לשיפור'},
    7: {1:'רמה 1 – הימנעות ובושה', 2:'רמה 2 – ציות פסיבי',
        3:'רמה 3 – סקרנות בסיסית', 4:'רמה 4 – חיפוש עזרה אקטיבי'},
    8: {1:'רמה 1 – חוסר מטרות', 2:'רמה 2 – שאלה טכנית/ציון',
        3:'רמה 3 – אבחון טעויות', 4:"רמה 4 – בקשת שיפור ומועד ב'"},
    9: {1:'רמה 1 – דחיינות מוחלטת', 2:'רמה 2 – תכנון ימים ספורים מראש',
        3:'רמה 3 – תכנון שבועי', 4:'רמה 4 – תכנון אסטרטגי ל-2 שבועות'},
}


def get_level_num(val):
    if val is None:
        return None
    m = re.search(r'רמה\s*(\d)', str(val))
    return int(m.group(1)) if m else None


def apply_correction(q, lv, ans):
    """Returns (new_level, reason) or (None, None) if no change."""
    if ans is None:
        return None, None
    a = str(ans).strip()

    if q == 1:
        # L2→L1: bare tool name, no explanation
        if lv == 2 and re.match(
            r'^(chat\s*gpt|chatgpt|צ[׳\'״]?אט\s*$|^gpt$|ג[׳\'״]?יפיטי|'
            r'^ai$|^בינה$|gemini|copilot|ישתמש\s*בצאט|אפתח\s*gpt|'
            r'^גוגל$|^יוטיוב$)[\s.,]*$',
            a, re.IGNORECASE
        ):
            return 1, 'מחוון L1: כלי בודד ללא הסבר כיצד משתמשים'
        # L2→L1: very short, no meaningful keywords
        if lv == 2 and len(a) <= 8 and not re.search(
            r'(AI|בינה|צאט|גוגל|יוטיוב|אבא|אמא|חבר|מורה)', a
        ):
            return 1, 'מחוון L1: תגובה ≤8 תווים ללא כוונה פעילה'
        # L2→L3: AI/search with no person dependency, len > 8
        if lv == 2 and len(a) > 8:
            person = re.search(r'(אבא|אמא|הורים|חבר|חברה|מורה|משפחה|אחי|אחות|סבא|סבתא)', a)
            indep  = re.search(
                r'(גוגל|יוטיוב|אינטרנט|AI|ai|בינה|צ[׳\'״]?אט|gpt|ג[׳\'״]?יפיטי|חפש|אחפש|סרטון)',
                a, re.IGNORECASE
            )
            if indep and not person:
                return 3, 'מחוון L3: AI/חיפוש עצמאי – יוזמה אקטיבית ללא הסתמכות על אדם'
        # L3→L4: independent search + self-testing
        if lv == 3 and re.search(
            r'(אבחן|בחן|חידון|מבחן|שאלון|אוודא|תרגול|אתרגל|תרגל|בדוק|בודק)', a
        ):
            return 4, 'מחוון L4: חיפוש עצמאי + בחינה עצמית מוצהרת'

    elif q == 2:
        # L2→L1: short, no attempt keywords
        if lv == 2 and len(a) <= 12 and not re.search(
            r'(אנסה|ינסה|אבדוק|אחפש|קרא|יוטיוב|גוגל)', a
        ):
            return 1, 'מחוון L1: פנייה מיידית ללא ניסיון עצמי קודם'
        # L2→L3: self-try + person + conditional/sequential
        if lv == 2:
            self_try = re.search(r'(קרא\s*שוב|חפש|יוטיוב|גוגל|אנסה|ינסה|שוב|מחדש)', a)
            person   = re.search(r'(אבא|אמא|הורים|חבר|מורה|משפחה)', a)
            after    = re.search(r'(אחרי|אם\s*לא|אם\s*עדיין|ולאחר\s*מכן|רק\s*אחר)', a)
            if self_try and person and after:
                return 3, 'מחוון L3: ניסיון עצמי → פנייה לאדם רק אחרי כישלון'

    elif q == 3:
        # L1→L2: internal attribution "לא למדתי"
        if lv == 1 and re.search(
            r'(לא\s*למדתי|לא\s*השקעתי|לא\s*הכנתי|לא\s*הבנתי|לא\s*תרגלתי|לא\s*ידעתי|לא\s*קראתי)', a
        ):
            return 2, 'מחוון L2: ייחוס פנימי כללי – הכרה ב"לא למדתי מספיק"'
        # L2→L1: external only – stress/blackout no responsibility
        if lv == 2 and re.match(
            r'^(הלחץ|בלאק[\s\-]?אאוט|חירטטתי|ככה\s*ככה|'
            r'כי\s*לא\s*עניתי\s*נכון|בלבול|לחץ|מהלחץ|'
            r'הייתי\s*בלחץ|בגלל\s*לחץ)[\s.,]*$',
            a.strip()
        ):
            return 1, 'מחוון L1: ייחוס חיצוני בלבד – לחץ/בלאק-אאוט ללא לקיחת אחריות'
        # L2→L3: identifies specific failure mechanism
        if lv == 2 and re.search(
            r'(שיטת\s*הלמידה|דרך\s*לא\s*נכונה|לא\s*תרגלתי|שאלות\s*מילוליות|'
            r'לא\s*חזרתי|הזנחתי|טעויות\s*קטנות\s*בחישוב|לא\s*קראתי\s*את\s*השאלה)', a
        ):
            return 3, 'מחוון L3: זיהוי מנגנון ספציפי של הכשל'

    elif q == 4:
        # L2→L3: internal subjective measure
        if lv == 2 and re.search(r'(היה\s*לי\s*קל|הרגשתי|הצלחתי\s*לענות|בזמן\s*המבחן|היה\s*קל)', a):
            if not re.search(r'(לפי\s*הציון\s*בלבד|הציון\s*הוא)', a):
                return 3, 'מחוון L3: מדד פנימי סובייקטיבי – "אם היה לי קל במבחן"'

    elif q == 6:
        action = (r'(אשתפר|ישפר|לשפר|ללמוד|אלמד|אנסה|ינסה|אשתדל|לנסות|'
                  r'אבוא|למבחן|יותר\s*טוב|השתפר|אשנה|אשקיע|ישקיע|'
                  r'לפעם\s*הבאה|הבאה|נסה|אטרח|אוסיף|אדבר)')
        if lv == 2 and len(a) <= 40 and not re.search(action, a):
            return 1, 'מחוון L1: תגובה רגשית ללא כוונת פעולה'

    elif q == 9:
        # L2→L1: very short, no time plan
        if lv == 2 and len(a) <= 15 and not re.search(r'(שבוע|יום|שעה|לפני|מראש)', a):
            return 1, 'מחוון L1: הצהרה כללית ללא תוכנית זמן'
        # L2→L3: specific time + internal readiness indicator
        if lv == 2 and re.search(r'(שבוע|שבועיים|יום|שעה|מראש|מוקדם)', a) and re.search(
            r'(אדע\s*שמוכן|ארגיש|מבחן\s*לדוגמ|אבחן|שאלון|תרגילים)', a
        ):
            return 3, 'מחוון L3: לוח זמנים ספציפי + מדד פנימי למוכנות'

    return None, None


def safe_write(ws, row, col, value):
    """Write to a cell only if it is not a MergedCell."""
    cell = ws.cell(row, col)
    if not isinstance(cell, MergedCell):
        cell.value = value


def process_sheet(wb, q):
    ws = wb[f'Q{q}']
    labels = LEVEL_LABELS[q]
    max_col = ws.max_column

    # ── locate key rows ──────────────────────────────────────────────────────
    data_header_row = None
    summary_rows = {}   # level_num → row index in summary table
    total_row = None
    avg_row   = None

    for r in range(1, 30):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value

        # data section column-header row
        if a == 'רמה' and b == 'עם/ללא AI':
            data_header_row = r

        # summary table level rows (col A = full level label, col C = count int/0)
        if (a and isinstance(a, str) and b and isinstance(b, str)
                and 'רמה' in a and get_level_num(a) is not None
                and b not in ('עם/ללא AI',) and isinstance(c, (int, float))):
            summary_rows[get_level_num(a)] = r

        if a == 'סה"כ' and not isinstance(ws.cell(r, 3), MergedCell):
            total_row = r
        if a == 'רמה ממוצעת':
            avg_row = r

    if data_header_row is None:
        print(f'  [שגיאה] Q{q}: data header row not found')
        return 0

    # ── read data rows ────────────────────────────────────────────────────────
    data_rows = []
    for r in range(data_header_row + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a is None or b is None:
            continue
        lv_n = get_level_num(a)
        if lv_n is None:
            continue
        if b in ('עם AI', 'ללא AI'):
            data_rows.append([
                lv_n,
                b,
                ws.cell(r, 3).value,   # school
                ws.cell(r, 4).value,   # answer
                ws.cell(r, 5).value,   # reason
            ])

    # ── apply corrections ─────────────────────────────────────────────────────
    corrections = 0
    for row in data_rows:
        lv, ai, school, ans, reason = row
        new_lv, new_reason = apply_correction(q, lv, ans)
        if new_lv is not None and new_lv != lv:
            row[0] = new_lv
            row[4] = new_reason
            corrections += 1

    if corrections == 0:
        return 0

    # ── re-sort: level asc, ללא AI before עם AI ──────────────────────────────
    data_rows.sort(key=lambda x: (x[0], 0 if x[1] == 'ללא AI' else 1))

    # ── recount ───────────────────────────────────────────────────────────────
    counts = defaultdict(lambda: {'עם': 0, 'ללא': 0})
    for lv, ai, *_ in data_rows:
        counts[lv]['עם' if ai == 'עם AI' else 'ללא'] += 1

    total     = len(data_rows)
    total_ai  = sum(v['עם'] for v in counts.values())
    total_no  = sum(v['ללא'] for v in counts.values())
    avg       = (sum(lv * (counts[lv]['עם'] + counts[lv]['ללא']) for lv in range(1, 5))
                 / total) if total else 0

    # ── update summary table ──────────────────────────────────────────────────
    for lv_n in range(1, 5):
        r = summary_rows.get(lv_n)
        if r is None:
            continue
        n_ai  = counts[lv_n]['עם']
        n_no  = counts[lv_n]['ללא']
        n_tot = n_ai + n_no
        safe_write(ws, r, 3, n_ai)
        safe_write(ws, r, 4, f'{n_ai / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, r, 5, n_no)
        safe_write(ws, r, 6, f'{n_no / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, r, 7, n_tot)
        safe_write(ws, r, 8, f'{n_tot / total * 100:.1f}%' if total else '0.0%')

    if total_row:
        safe_write(ws, total_row, 3, total_ai)
        safe_write(ws, total_row, 4, f'{total_ai / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, total_row, 5, total_no)
        safe_write(ws, total_row, 6, f'{total_no / total * 100:.1f}%' if total else '0.0%')
        safe_write(ws, total_row, 7, total)
        safe_write(ws, total_row, 8, '100%')

    if avg_row:
        safe_write(ws, avg_row, 2, f'{avg:.2f}')

    # update info line (row just before data_header_row)
    info_row = data_header_row - 1
    a_info = ws.cell(info_row, 1).value
    if a_info and 'כל התשובות' in str(a_info):
        safe_write(ws, info_row, 6, f'סה"כ: {total} | AI: {total_ai}')

    # ── rewrite raw data section ──────────────────────────────────────────────
    write_r = data_header_row + 1
    current_level = None

    for lv, ai, school, ans, reason in data_rows:
        if lv != current_level:
            current_level = lv
            n_in_level = counts[lv]['עם'] + counts[lv]['ללא']
            safe_write(ws, write_r, 1, f'{labels[lv]} ({n_in_level} תשובות)')
            for c in range(2, max_col + 1):
                safe_write(ws, write_r, c, None)
            write_r += 1

        safe_write(ws, write_r, 1, f'רמה {lv}')
        safe_write(ws, write_r, 2, ai)
        safe_write(ws, write_r, 3, school)
        safe_write(ws, write_r, 4, ans)
        safe_write(ws, write_r, 5, reason)
        for c in range(6, max_col + 1):
            safe_write(ws, write_r, c, None)
        write_r += 1

    # clear leftover rows
    while write_r <= ws.max_row:
        for c in range(1, max_col + 1):
            safe_write(ws, write_r, c, None)
        write_r += 1

    return corrections


# ─── main ────────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(INPUT_PATH)
grand_total = 0

for q in range(1, 10):
    n = process_sheet(wb, q)
    if n:
        print(f'  ✓ Q{q}: {n} תיקונים')
        grand_total += n
    else:
        print(f'    Q{q}: אין תיקונים')

wb.save(OUTPUT_PATH)
print(f'\n✓ נשמר: {OUTPUT_PATH}')
print(f'סה"כ תיקונים: {grand_total}')

"""
תיקון מיון תשובות גולמיות בגיליונות תשובות Q1–Q9
מיון: ראשית לפי רמה (1→5), בתוך כל רמה — ללא AI לפני עם AI
"""

import openpyxl
import re

# ─── נתיבים ────────────────────────────────────────────────────────────────
INPUT_PATH  = r"D:\קבוצת המחקר עם מנוחה\מאמר 3\קובץ עדכני\dashboard_מלא_עם_ביקורת (4).xlsx"
OUTPUT_PATH = r"D:\קבוצת המחקר עם מנוחה\מאמר 3\קובץ עדכני\dashboard_מיון_מתוקן.xlsx"
# ───────────────────────────────────────────────────────────────────────────


def level_num(val):
    """מחלץ מספר רמה מתוך טקסט כגון 'רמה 3' או 'רמה 3 – ...' """
    if val is None:
        return 99
    m = re.search(r'רמה\s*(\d)', str(val))
    return int(m.group(1)) if m else 99


def is_section_header(val):
    """שורת כותרת ביניים — 'רמה X – ...' """
    if val is None:
        return False
    return bool(re.search(r'רמה\s*\d\s*[–—\-]', str(val)))


def ai_order(val):
    """ללא AI=0 (ראשון), עם AI=1 (שני)"""
    if val is None:
        return 99
    s = str(val)
    if 'ללא' in s:
        return 0
    if 'עם' in s:
        return 1
    return 99


def row_is_empty(row_vals):
    return not any(v is not None and str(v).strip() for v in row_vals)


wb = openpyxl.load_workbook(INPUT_PATH)
fixed = []

for q in range(1, 10):
    sheet_name = f"תשובות Q{q}"
    if sheet_name not in wb.sheetnames:
        print(f"  [דלג] גיליון לא נמצא: {sheet_name}")
        continue

    ws = wb[sheet_name]
    max_col = ws.max_column

    # ── מצא את שורת כותרות העמודות (השורה שמכילה "רמה" כערך מדויק) ──
    header_row = None
    for r in range(1, min(20, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, max_col + 1)]
        if 'רמה' in vals:
            header_row = r
            break

    if header_row is None:
        print(f"  [שגיאה] לא נמצאה שורת כותרות ב-{sheet_name}")
        continue

    header_vals = [ws.cell(header_row, c).value for c in range(1, max_col + 1)]

    # ── מצא אינדקסי עמודות (1-based) ──
    try:
        col_level = next(i + 1 for i, v in enumerate(header_vals) if v == 'רמה')
        col_ai    = next(i + 1 for i, v in enumerate(header_vals) if v == 'עם/ללא AI')
    except StopIteration:
        print(f"  [שגיאה] עמודות רמה/AI לא נמצאו ב-{sheet_name}")
        continue

    # ── קרא שורות נתונים ושורות כותרת-ביניים ──
    data_rows    = []   # (level_num, ai_order, [values...])
    sec_headers  = {}   # level_num -> [values...]  (כותרת ביניים לכל רמה)

    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, max_col + 1)]
        if row_is_empty(vals):
            continue
        lv = vals[col_level - 1]
        if is_section_header(lv):
            n = level_num(lv)
            if n not in sec_headers:
                sec_headers[n] = vals
        else:
            data_rows.append((level_num(lv), ai_order(vals[col_ai - 1]), vals))

    # ── מיין: רמה עולה → ללא AI לפני עם AI ──
    data_rows.sort(key=lambda x: (x[0], x[1]))

    # ── כתוב חזרה לגיליון ──
    write_r = header_row + 1
    current_level = None

    for lv, ai, vals in data_rows:
        # כתוב כותרת ביניים כשמשתנה הרמה
        if lv != current_level:
            current_level = lv
            if lv in sec_headers:
                for c, v in enumerate(sec_headers[lv], 1):
                    ws.cell(write_r, c, v)
                write_r += 1

        # כתוב שורת נתון
        for c, v in enumerate(vals, 1):
            ws.cell(write_r, c, v)
        write_r += 1

    # מחק שורות עודפות שנשארו מהמבנה הישן
    while write_r <= ws.max_row:
        for c in range(1, max_col + 1):
            ws.cell(write_r, c, None)
        write_r += 1

    fixed.append(sheet_name)
    print(f"  ✓ תוקן: {sheet_name}  ({len(data_rows)} תשובות)")

wb.save(OUTPUT_PATH)
print(f"\n✓ הקובץ המתוקן נשמר:\n  {OUTPUT_PATH}")
print(f"גיליונות שתוקנו: {', '.join(fixed)}")

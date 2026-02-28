"""
טבלת התפלגות לפי שאלה ורמה – עם הבחנה בין שימוש ב-AI וללא שימוש
N = 1,250 תלמידי כיתה ז | מחוון פעלנות לומדים
"""

import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import random

random.seed(2024)
np.random.seed(2024)

# ─────────────────────────────────────────────
# מחוון (זהה למחוון_מפורט_סימולציה)
# ─────────────────────────────────────────────
RUBRIC = {
    "motivation":      {"he": "א. מוטיבציה ורלוונטיות", "questions": ["Q50","Q51"]},
    "growth_mindset":  {"he": "ב. תודעת צמיחה",          "questions": ["Q52","Q53"]},
    "initiative":      {"he": "ג. יוזמה ואחריות",         "questions": ["Q54"]},
    "self_regulation": {"he": "ד. ויסות עצמי",            "questions": ["Q55","Q56"]},
    "self_awareness":  {"he": "ה. מודעות עצמית",          "questions": ["Q57","Q58"]},
    "support":         {"he": "ו. תמיכה וחוויות רגשיות",  "questions": ["Q59"]},
}
DIM_CODES = list(RUBRIC.keys())

LEVEL_NAMES = {1:"בתחילת הדרך", 2:"מתפתח/ת", 3:"מתקדם/ת", 4:"מיומן/ת", 5:"מומחה/ית"}

Q_TEXT = {
    "Q50": "Q50 – רלוונטיות הלמידה לחיים",
    "Q51": "Q51 – מוכנות להתמודד עם אתגרים",
    "Q52": "Q52 – אמונה ביכולת השתפרות",
    "Q53": "Q53 – ראיית קושי כהזדמנות",
    "Q54": "Q54 – יוזמה ואחריות עצמאית",
    "Q55": "Q55 – תכנון וניהול זמן",
    "Q56": "Q56 – ויסות רגשי תחת לחץ",
    "Q57": "Q57 – היכרות עם חוזקות/חולשות",
    "Q58": "Q58 – הפקת לקחים",
    "Q59": "Q59 – פנייה לעזרה ותמיכה",
}

Q_DIM = {
    "Q50":"motivation","Q51":"motivation",
    "Q52":"growth_mindset","Q53":"growth_mindset",
    "Q54":"initiative",
    "Q55":"self_regulation","Q56":"self_regulation",
    "Q57":"self_awareness","Q58":"self_awareness",
    "Q59":"support",
}

QUESTIONS = ["Q50","Q51","Q52","Q53","Q54","Q55","Q56","Q57","Q58","Q59"]
AI_TOOLS  = ["ChatGPT","Claude","Gemini","Google Bard","Copilot","לא משתמש/ת"]
AI_W      = [0.30, 0.12, 0.18, 0.10, 0.08, 0.22]
N         = 1250

DIM_DIST  = {
    "motivation":      [0.06,0.16,0.38,0.28,0.12],
    "growth_mindset":  [0.12,0.22,0.36,0.22,0.08],
    "initiative":      [0.05,0.13,0.32,0.32,0.18],
    "self_regulation": [0.08,0.17,0.36,0.27,0.12],
    "self_awareness":  [0.09,0.19,0.38,0.25,0.09],
    "support":         [0.11,0.20,0.37,0.22,0.10],
}

def assign_level(s):
    if s < 1.5: return 1
    if s < 2.5: return 2
    if s < 3.5: return 3
    if s < 4.5: return 4
    return 5

# ─────────────────────────────────────────────
# יצירת נתונים
# ─────────────────────────────────────────────
def generate_students(n):
    rows = []
    for sid in range(1001, 1001+n):
        ai_tool = np.random.choice(AI_TOOLS, p=AI_W)
        ai_user = ai_tool != "לא משתמש/ת"
        dim_scores = {}
        for dim in DIM_CODES:
            probs = DIM_DIST[dim].copy()
            if ai_user and dim in ("initiative","self_regulation"):
                probs[2] = max(0, probs[2]-0.02)
                probs[2] -= 0.02
                probs[3] += 0.03; probs[4] += 0.03
                s = sum(probs); probs = [p/s for p in probs]
            lvl = np.random.choice([1,2,3,4,5], p=probs)
            score = np.clip(np.random.uniform(lvl-0.49, lvl+0.49), 1.0, 5.0)
            dim_scores[dim] = round(score, 1)

        q_ans = {}
        for q, dim in Q_DIM.items():
            base  = dim_scores[dim]
            noise = random.choice([-0.5,0,0,0,0.5])
            q_ans[q] = int(np.clip(round(base+noise), 1, 5))

        row = {"מזהה":sid, "כיתה":"ז", "AI_tool":ai_tool, "ai_user":ai_user}
        row.update(q_ans)
        rows.append(row)
    return pd.DataFrame(rows)

print("⏳ יוצר נתונים...")
df = generate_students(N)
ai_df    = df[df["ai_user"]==True]
no_ai_df = df[df["ai_user"]==False]
print(f"✅  סה\"כ: {len(df)} | עם AI: {len(ai_df)} | ללא AI: {len(no_ai_df)}")

# ─────────────────────────────────────────────
# חישוב אחוזים לכל שאלה × רמה × קבוצה
# ─────────────────────────────────────────────
def pct_dist(sub, q):
    c = sub[q].value_counts()
    total = len(sub)
    return {lvl: round(c.get(lvl,0)/total*100, 1) for lvl in range(1,6)}

stats = {}
for q in QUESTIONS:
    stats[q] = {
        "all":   pct_dist(df, q),
        "ai":    pct_dist(ai_df, q),
        "no_ai": pct_dist(no_ai_df, q),
    }

# ─────────────────────────────────────────────
# צבעים – זהים לדוגמה
# ─────────────────────────────────────────────
COL_HEADER  = "1C2D5E"   # כחול כהה
COL_Q_BG    = "2C3E7B"   # כחול בינוני לשם השאלה
COL_Q_FONT  = "FFFFFF"

# רמות – בדיוק כמו בתמונה
LVL_BG = {1:"F4CCCC", 2:"FCE5CD", 3:"FFF2CC", 4:"D9EAD3", 5:"D0E0F0"}
LVL_FONT_BOLD = {1:"CC0000", 2:"E65C00", 3:"7A6000", 4:"274E13", 5:"1A4875"}

GROUP_BG = {
    "no_ai": "EEF2FF",   # גוון כחול-לבן לשורת ללא AI
    "ai":    "F0FFF0",   # גוון ירוק-לבן לשורת עם AI
}
GROUP_FONT = {"no_ai":"1C2D5E", "ai":"1A5C1A"}

# ─────────────────────────────────────────────
# כלי עזר לעיצוב תאים
# ─────────────────────────────────────────────
def rtl(text):
    return any('\u0590' <= c <= '\u05FF' for c in str(text))

def cs(ws, r, c, val, bold=False, sz=10, fc="000000", bg=None,
       wrap=False, ah="center", av="center", mr=None, italic=False):
    cell = ws.cell(row=r, column=c, value=val)
    cell.font = Font(bold=bold, size=sz, color=fc, italic=italic,
                     name="David" if rtl(str(val)) else "Calibri")
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal=ah, vertical=av, wrap_text=wrap)
    thin = Side(style="thin", color="BDBDBD")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if mr:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=mr)
    return cell

# ─────────────────────────────────────────────
# בנה Excel
# ─────────────────────────────────────────────
wb = openpyxl.Workbook()
wb.remove(wb.active)

# ════════════════════════════════════════════
# גיליון 1 – טבלת אחוזים (עיצוב כמו בדוגמה)
# ════════════════════════════════════════════
ws = wb.create_sheet("התפלגות AI vs ללא AI")
ws.sheet_view.rightToLeft = True

# ── כותרת ראשית ──────────────────────────────
cs(ws,1,1, "סיכום סטטיסטי – התפלגות תלמידים לפי שאלה ורמה",
   bold=True, sz=14, fc="FFFFFF", bg=COL_HEADER, mr=8)
ws.row_dimensions[1].height = 28

cs(ws,2,1,
   f"N={N:,} תלמידי כיתה ז  |  הבחנה: ✦ ללא AI (n={len(no_ai_df)}) vs  ✦ עם AI (n={len(ai_df)})",
   bold=False, sz=10, fc="FFFFFF", bg="37474F", mr=8)
ws.row_dimensions[2].height = 18

# ── כותרות עמודות ────────────────────────────
r = 4
# שורת כותרות רמות
cs(ws,r,1, "שאלה",   bold=True, sz=11, fc=COL_Q_FONT, bg=COL_HEADER, mr=2)
cs(ws,r,3, "קבוצה",  bold=True, sz=10, fc=COL_Q_FONT, bg=COL_HEADER)
for lvl in range(1,6):
    cs(ws,r, 3+lvl,
       f"רמה {lvl}\n{LEVEL_NAMES[lvl]}",
       bold=True, sz=10, fc=LVL_FONT_BOLD[lvl], bg=LVL_BG[lvl], wrap=True)
ws.row_dimensions[r].height = 36

# ── שורות נתונים ─────────────────────────────
r = 5
for q in QUESTIONS:
    q_label = Q_TEXT[q]
    dim_he  = RUBRIC[Q_DIM[q]]["he"]

    # שורת שם השאלה + ממד (ממוזגת)
    cs(ws, r, 1,
       f"{q_label}\n({dim_he})",
       bold=True, sz=10, fc=COL_Q_FONT, bg=COL_Q_BG,
       wrap=True, ah="right", mr=2)
    ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=2)

    # שתי שורות: ללא AI / עם AI
    for g_key, g_label, g_bg, g_fc in [
        ("no_ai", "ללא AI ✧", GROUP_BG["no_ai"], GROUP_FONT["no_ai"]),
        ("ai",    "עם AI  ✦", GROUP_BG["ai"],    GROUP_FONT["ai"]),
    ]:
        cs(ws, r, 3, g_label, bold=True, sz=9, fc=g_fc, bg=g_bg)
        for lvl in range(1,6):
            pct = stats[q][g_key][lvl]
            cs(ws, r, 3+lvl,
               f"{pct:.1f}%",
               bold=(pct >= 35), sz=10,
               fc=LVL_FONT_BOLD[lvl], bg=LVL_BG[lvl])
        ws.row_dimensions[r].height = 22
        r += 1

    # קו הפרדה (שורה ריקה קלה)
    for col in range(1, 9):
        cs(ws, r, col, "", bg="E8EBF0")
    ws.row_dimensions[r].height = 5
    r += 1

# ── רוחבי עמודות ────────────────────────────
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 2
ws.column_dimensions["C"].width = 12
for letter in ["D","E","F","G","H"]:
    ws.column_dimensions[letter].width = 13

# ════════════════════════════════════════════
# גיליון 2 – כולל N בכל תא (לא רק %)
# ════════════════════════════════════════════
ws2 = wb.create_sheet("N ואחוז לפי שאלה")
ws2.sheet_view.rightToLeft = True

cs(ws2,1,1, "N ואחוז – התפלגות לפי שאלה, רמה וקבוצה",
   bold=True, sz=13, fc="FFFFFF", bg=COL_HEADER, mr=8)
ws2.row_dimensions[1].height = 26

r = 3
cs(ws2,r,1, "שאלה",  bold=True, sz=11, fc=COL_Q_FONT, bg=COL_HEADER, mr=2)
cs(ws2,r,3, "קבוצה", bold=True, sz=10, fc=COL_Q_FONT, bg=COL_HEADER)
for lvl in range(1,6):
    cs(ws2,r, 3+lvl,
       f"רמה {lvl}\n{LEVEL_NAMES[lvl]}",
       bold=True, sz=10, fc=LVL_FONT_BOLD[lvl], bg=LVL_BG[lvl], wrap=True)
ws2.row_dimensions[r].height = 36

r = 4
for q in QUESTIONS:
    q_label = Q_TEXT[q]
    dim_he  = RUBRIC[Q_DIM[q]]["he"]

    cs(ws2, r, 1,
       f"{q_label}\n({dim_he})",
       bold=True, sz=10, fc=COL_Q_FONT, bg=COL_Q_BG,
       wrap=True, ah="right", mr=2)
    ws2.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=2)

    for g_key, g_label, n_group, g_bg, g_fc in [
        ("no_ai", "ללא AI ✧", len(no_ai_df), GROUP_BG["no_ai"], GROUP_FONT["no_ai"]),
        ("ai",    "עם AI  ✦", len(ai_df),    GROUP_BG["ai"],    GROUP_FONT["ai"]),
    ]:
        cs(ws2, r, 3, f"{g_label}\n(n={n_group})", bold=True, sz=9,
           fc=g_fc, bg=g_bg, wrap=True)
        for lvl in range(1,6):
            pct = stats[q][g_key][lvl]
            cnt = round(pct / 100 * n_group)
            cs(ws2, r, 3+lvl,
               f"n={cnt}\n{pct:.1f}%",
               bold=(pct >= 35), sz=9,
               fc=LVL_FONT_BOLD[lvl], bg=LVL_BG[lvl], wrap=True)
        ws2.row_dimensions[r].height = 30
        r += 1

    for col in range(1,9):
        cs(ws2, r, col, "", bg="E8EBF0")
    ws2.row_dimensions[r].height = 4
    r += 1

ws2.column_dimensions["A"].width = 26
ws2.column_dimensions["B"].width = 2
ws2.column_dimensions["C"].width = 14
for letter in ["D","E","F","G","H"]:
    ws2.column_dimensions[letter].width = 13

# ════════════════════════════════════════════
# גיליון 3 – כלל (ללא הבחנה) – כמו בדוגמה המקורית
# ════════════════════════════════════════════
ws3 = wb.create_sheet("כלל האוכלוסייה")
ws3.sheet_view.rightToLeft = True

cs(ws3,1,1, "סיכום סטטיסטי – התפלגות תלמידים לפי שאלה ורמה (כלל האוכלוסייה)",
   bold=True, sz=13, fc="FFFFFF", bg=COL_HEADER, mr=7)
ws3.row_dimensions[1].height = 26

r = 3
cs(ws3,r,1, "שאלה", bold=True, sz=11, fc=COL_Q_FONT, bg=COL_HEADER)
for lvl in range(1,6):
    cs(ws3,r, 1+lvl,
       f"רמה {lvl}\n{LEVEL_NAMES[lvl]}",
       bold=True, sz=10, fc=LVL_FONT_BOLD[lvl], bg=LVL_BG[lvl], wrap=True)
cs(ws3,r,7, "ממוצע", bold=True, sz=10, fc=COL_Q_FONT, bg=COL_HEADER)
ws3.row_dimensions[r].height = 36
r = 4

for q in QUESTIONS:
    q_label = Q_TEXT[q]
    dim_he  = RUBRIC[Q_DIM[q]]["he"]
    cs(ws3,r,1, f"{q_label}  ({dim_he})",
       bold=True, sz=10, fc=COL_Q_FONT, bg=COL_Q_BG, ah="right")
    for lvl in range(1,6):
        pct = stats[q]["all"][lvl]
        cs(ws3,r, 1+lvl,
           f"{pct:.1f}%",
           bold=(pct >= 35), sz=11,
           fc=LVL_FONT_BOLD[lvl], bg=LVL_BG[lvl])
    mean_q = df[q].mean()
    cs(ws3,r,7, f"{mean_q:.2f}", bold=True, sz=10, bg="ECEFF1")
    ws3.row_dimensions[r].height = 22
    r += 1

ws3.column_dimensions["A"].width = 36
for letter in ["B","C","D","E","F"]:
    ws3.column_dimensions[letter].width = 14
ws3.column_dimensions["G"].width = 10

# ── שמירה ────────────────────────────────────
OUT = "התפלגות_AI_vs_ללא_AI.xlsx"
wb.save(OUT)
print(f"\n✅ נשמר: {OUT}")

# ── הדפסת סיכום ──────────────────────────────
print(f"\n{'='*70}")
print(f"הבחנה: ללא AI (n={len(no_ai_df)}) | עם AI (n={len(ai_df)})")
print(f"{'='*70}")
print(f"{'שאלה':<7} {'קבוצה':<9}  " +
      "  ".join([f"{'רמה'+str(l):<9}" for l in range(1,6)]))
print("-"*70)
for q in QUESTIONS:
    for g_key, g_label in [("no_ai","ללא AI"), ("ai","עם AI ")]:
        parts = [f"{q:<7}", f"{g_label:<9}"]
        for lvl in range(1,6):
            pct = stats[q][g_key][lvl]
            parts.append(f"{pct:5.1f}%   ")
        print("  ".join(parts))
    print()
